"""Add Backend Progress Tracker + % summary to Project Schedule.xlsx (match Frontend layout)."""

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PATH = r"C:\xampp\htdocs\leaveforms-api\docs\Project Schedule.xlsx"

# Section, Task, Status — Status: Completed | In Progress | Not Started | Blocked
# Progress % formula will derive from Status (same scoring as FE: Completed=1, In Progress=0.5)
TRACKER_ROWS = [
    # 1. Foundation
    ("1. Foundation", "Dual DB setup — formsdb + kdtphdb (read-only)", "Completed"),
    ("1. Foundation", "Portal cookie auth (userID → kdtlogin)", "Completed"),
    ("1. Foundation", "Routes — GET /api/me, GET /api/health", "Completed"),
    ("1. Foundation", "Schema migrations — 8 revamp tables + indexes", "Completed"),
    ("1. Foundation", "Eloquent models + PHP backed enums", "Completed"),
    ("1. Foundation", "LeaveTypeSeeder (VL, SL, EL, ML, PL, PTB, LL)", "Completed"),
    ("1. Foundation", "Config — president_emp_id, CORS, attachments, notifications", "Completed"),
    ("1. Foundation", "Staging/prod portal cookie & env verify with infra", "In Progress"),
    # 2. Employee Features
    ("2. Employee Features", "POST /leaves — submit VL with validation", "Completed"),
    ("2. Employee Features", "Duration calc — workdays, holidays, half-day", "Completed"),
    ("2. Employee Features", "paid_days / lwop_days split + reserve credits", "Completed"),
    ("2. Employee Features", "Snapshot approval requirements at submit", "Completed"),
    ("2. Employee Features", "File on behalf (filed_by_emp_id / Group Leader)", "Completed"),
    ("2. Employee Features", "Attachments — store / replace / download", "Completed"),
    ("2. Employee Features", "GET /leaves — list / LeaveResource", "Completed"),
    ("2. Employee Features", "PUT/POST /leaves/{id} — modify while pending", "Completed"),
    ("2. Employee Features", "POST /leaves/{id}/cancel", "Completed"),
    ("2. Employee Features", "GET /balances — read leave credits", "Completed"),
    ("2. Employee Features", "Email notify approvers on submit", "Completed"),
    ("2. Employee Features", "Reject notify to employee / filer", "Not Started"),
    ("2. Employee Features", "Notify next-tier approvers after partial approve", "Not Started"),
    ("2. Employee Features", "Submit flows for other leave types (EL/ML/PL/PTB/LL)", "Not Started"),
    # 3. Approver Features
    ("3. Approver Features", "GET /approvals/pending — formspic group scope", "Completed"),
    ("3. Approver Features", "Approve path — partially_approved → confirmed", "Completed"),
    ("3. Approver Features", "Reject path — close leave + release reserve", "Completed"),
    ("3. Approver Features", "Approval policy matrix (regular → P4 / president)", "Completed"),
    ("3. Approver Features", "Finalize — confirmPaidDays / skip ledger if full LWOP", "Completed"),
    ("3. Approver Features", "Notify employee on confirm", "Completed"),
    ("3. Approver Features", "President fallback verify (staging emp ID)", "In Progress"),
    ("3. Approver Features", "Stakeholder — special group rules (KRM / checkTier)", "Not Started"),
    ("3. Approver Features", "Admin-requirement approval paths for other leave types", "Not Started"),
    # 4. Admin Features
    ("4. Admin Features", "leave_conversion schema + model + enums", "Completed"),
    ("4. Admin Features", "Opening balances policy — manual entry (no legacy leave import)", "Completed"),
    ("4. Admin Features", "applyInitialBalance + applyManualAdjustment (service)", "Completed"),
    ("4. Admin Features", "Admin/HR API — set opening balance (manual from legacy)", "Not Started"),
    ("4. Admin Features", "Admin/HR API — manual adjust with remarks", "Not Started"),
    ("4. Admin Features", "VL → SL conversion request + med cert upload", "Not Started"),
    ("4. Admin Features", "Admin conversion review — approve / reject", "Not Started"),
    ("4. Admin Features", "Conversion balance — vl_to_sl_restore + deduct", "Not Started"),
    ("4. Admin Features", "Conversion notify outcome", "Not Started"),
    ("4. Admin Features", "applyAnnualEntitlement / monthly accrual / jobs", "Not Started"),
    ("4. Admin Features", "reverseConfirmedLeave (admin undo)", "Not Started"),
    ("4. Admin Features", "Report API — sum confirmed leave_info", "Not Started"),
    ("4. Admin Features", "Quarterly absence warnings", "Not Started"),
    ("4. Admin Features", "Other leave types — admin requirement rows + stakeholder sets", "Not Started"),
    # 5. Frontend Integration
    ("5. Frontend Integration (TBD)", "Wire FE screens to real API routes", "In Progress"),
    ("5. Frontend Integration (TBD)", "Portal cookie / CORS / env per environment", "In Progress"),
    ("5. Frontend Integration (TBD)", "Contract alignment — LeaveResource / balances payloads", "In Progress"),
    ("5. Frontend Integration (TBD)", "End-to-end smoke: submit → approve → confirm", "Not Started"),
    # 6. Testing (after development)
    ("6. Testing (after development)", "Feature tests — submit / approve / reject / cancel", "Not Started"),
    ("6. Testing (after development)", "Balance ledger invariant tests", "Not Started"),
    ("6. Testing (after development)", "Approval policy matrix tests", "Not Started"),
    ("6. Testing (after development)", "Conversion + notification tests", "Not Started"),
    # 7. Documentation (after development)
    ("7. Documentation (after development)", "API endpoint notes for FE partner", "Not Started"),
    ("7. Documentation (after development)", "Manual opening-balance process notes for HR", "Not Started"),
    ("7. Documentation (after development)", "Env / portal cookie cutover checklist", "Not Started"),
]

SECTIONS = [
    ("1. Foundation", "FFFCE4D6"),
    ("2. Employee Features", "FFDDEBF7"),
    ("3. Approver Features", "FFBDD7EE"),
    ("4. Admin Features", "FFE2D5F1"),
    ("5. Frontend Integration (TBD)", "FFF4CCCC"),
    ("6. Testing (after development)", "FFC6EFCE"),
    ("7. Documentation (after development)", "FFD9EAD3"),
]

STATUS_FILL = {
    "Completed": PatternFill("solid", fgColor="FFC6EFCE"),
    "In Progress": PatternFill("solid", fgColor="FFFFEB9C"),
    "Not Started": PatternFill("solid", fgColor="FFFFC7CE"),
    "Blocked": PatternFill("solid", fgColor="FFD9D9D9"),
}

THIN = Border(
    left=Side(style="thin", color="FFB0B0B0"),
    right=Side(style="thin", color="FFB0B0B0"),
    top=Side(style="thin", color="FFB0B0B0"),
    bottom=Side(style="thin", color="FFB0B0B0"),
)


def status_to_pct_formula(row: int) -> str:
    # Same scoring as FE Progress Tracker: Completed=1, In Progress=0.5, else 0
    return (
        f'=IF(C{row}="Completed",1,IF(C{row}="In Progress",0.5,0))'
    )


def main():
    wb = load_workbook(PATH)

    # --- Backend Progress Tracker sheet ---
    if "Backend Progress Tracker" in wb.sheetnames:
        del wb["Backend Progress Tracker"]
    tracker = wb.create_sheet("Backend Progress Tracker", 2)

    headers = ["Section", "Task", "Status", "Progress %"]
    header_fill = PatternFill("solid", fgColor="FF1F4E79")
    header_font = Font(bold=True, color="FFFFFFFF", size=11)
    for col, h in enumerate(headers, 1):
        cell = tracker.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN

    section_fills = {name: PatternFill("solid", fgColor=color) for name, color in SECTIONS}

    for i, (section, task, status) in enumerate(TRACKER_ROWS, start=2):
        tracker.cell(i, 1, section).fill = section_fills.get(section, PatternFill())
        tracker.cell(i, 1).font = Font(bold=True, size=10)
        tracker.cell(i, 2, task)
        st = tracker.cell(i, 3, status)
        st.fill = STATUS_FILL.get(status, PatternFill())
        st.alignment = Alignment(horizontal="center")
        pct = tracker.cell(i, 4, status_to_pct_formula(i))
        pct.number_format = "0%"
        pct.alignment = Alignment(horizontal="center")
        for c in range(1, 5):
            tracker.cell(i, c).border = THIN
            tracker.cell(i, c).alignment = Alignment(
                vertical="center", wrap_text=True, horizontal=tracker.cell(i, c).alignment.horizontal
            )

    last_data = 1 + len(TRACKER_ROWS)

    # Data validation for Status
    dv = DataValidation(
        type="list",
        formula1='"Completed,In Progress,Not Started,Blocked"',
        allow_blank=False,
    )
    dv.error = "Pick Completed, In Progress, Not Started, or Blocked"
    dv.errorTitle = "Invalid status"
    tracker.add_data_validation(dv)
    dv.add(f"C2:C{last_data}")

    tracker.column_dimensions["A"].width = 32
    tracker.column_dimensions["B"].width = 70
    tracker.column_dimensions["C"].width = 14
    tracker.column_dimensions["D"].width = 12
    tracker.auto_filter.ref = f"A1:D{last_data}"
    tracker.freeze_panes = "A2"

    note_row = last_data + 2
    tracker.cell(
        note_row,
        1,
        "Commitment: backend DEVELOPMENT by end of August (weekdays). "
        "Phases 2–3 are Partial (VL only) until other leave types are done. "
        "Testing & Documentation after development (from Sep 2026). "
        "Legacy leaves stay in legacy app; opening balances entered manually. "
        "Change Status in column C — Progress % and the Backend Tentative Schedule summary update automatically.",
    ).font = Font(italic=True, color="FF666666", size=9)
    tracker.merge_cells(start_row=note_row, start_column=1, end_row=note_row, end_column=4)

    # --- Backend Tentative Schedule % block (mirror Frontend rows 19–30) ---
    be = wb["Backend Tentative Schedule"]

    # Clear any old content from row 18 downward that might conflict
    for row in be.iter_rows(min_row=18, max_row=40, max_col=7):
        for cell in row:
            cell.value = None
            cell.fill = PatternFill()
            cell.font = Font()
            cell.border = Border()

    # Labels
    be["A19"] = "Overall %"
    be["A19"].font = Font(bold=True, size=11, color="FF1F4E79")
    be["A19"].alignment = Alignment(horizontal="center")

    be["C19"] = "Completed"
    be["C19"].font = Font(bold=True, size=11)
    be["C19"].fill = PatternFill("solid", fgColor="FFC6EFCE")
    be["C19"].alignment = Alignment(horizontal="center")

    be["E19"] = "In Progress"
    be["E19"].font = Font(bold=True, size=11)
    be["E19"].fill = PatternFill("solid", fgColor="FFFFEB9C")
    be["E19"].alignment = Alignment(horizontal="center")

    be["G19"] = "Not Started"
    be["G19"].font = Font(bold=True, size=11)
    be["G19"].fill = PatternFill("solid", fgColor="FFFFC7CE")
    be["G19"].alignment = Alignment(horizontal="center")

    # Overall formulas — same pattern as FE but local sheet
    be["A20"] = '=IF(COUNTA(\'Backend Progress Tracker\'!C:C)<=1,0,AVERAGE(\'Backend Progress Tracker\'!D:D))'
    be["A20"].font = Font(bold=True, size=28, color="FF1F4E79")
    be["A20"].fill = PatternFill("solid", fgColor="FFD6EAF8")
    be["A20"].number_format = "0%"
    be["A20"].alignment = Alignment(horizontal="center", vertical="center")
    be["A20"].border = THIN

    be["C20"] = '=COUNTIF(\'Backend Progress Tracker\'!C:C,"Completed")'
    be["C20"].font = Font(bold=True, size=22)
    be["C20"].fill = PatternFill("solid", fgColor="FFC6EFCE")
    be["C20"].alignment = Alignment(horizontal="center", vertical="center")
    be["C20"].border = THIN

    be["E20"] = '=COUNTIF(\'Backend Progress Tracker\'!C:C,"In Progress")'
    be["E20"].font = Font(bold=True, size=22)
    be["E20"].fill = PatternFill("solid", fgColor="FFFFEB9C")
    be["E20"].alignment = Alignment(horizontal="center", vertical="center")
    be["E20"].border = THIN

    be["G20"] = '=COUNTIF(\'Backend Progress Tracker\'!C:C,"Not Started")'
    be["G20"].font = Font(bold=True, size=22)
    be["G20"].fill = PatternFill("solid", fgColor="FFFFC7CE")
    be["G20"].alignment = Alignment(horizontal="center", vertical="center")
    be["G20"].border = THIN

    be.row_dimensions[20].height = 36

    # Section table header
    section_headers = ["Section", "Tasks", "Completed", "In Progress", "Not Started", "Blocked", "Progress %"]
    for col, h in enumerate(section_headers, 1):
        cell = be.cell(22, col, h)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = THIN

    for i, (section, color) in enumerate(SECTIONS):
        row = 23 + i
        fill = PatternFill("solid", fgColor=color)
        be.cell(row, 1, section).fill = fill
        be.cell(row, 1).font = Font(bold=True, size=11)
        # Tasks
        be.cell(row, 2, f"=COUNTIF('Backend Progress Tracker'!A:A,A{row})")
        # Completed
        be.cell(
            row,
            3,
            f"=COUNTIFS('Backend Progress Tracker'!A:A,A{row},'Backend Progress Tracker'!C:C,\"Completed\")",
        )
        # In Progress
        be.cell(
            row,
            4,
            f"=COUNTIFS('Backend Progress Tracker'!A:A,A{row},'Backend Progress Tracker'!C:C,\"In Progress\")",
        )
        # Not Started
        be.cell(
            row,
            5,
            f"=COUNTIFS('Backend Progress Tracker'!A:A,A{row},'Backend Progress Tracker'!C:C,\"Not Started\")",
        )
        # Blocked
        be.cell(
            row,
            6,
            f"=COUNTIFS('Backend Progress Tracker'!A:A,A{row},'Backend Progress Tracker'!C:C,\"Blocked\")",
        )
        # Progress % = (Completed + 0.5*In Progress) / Tasks
        be.cell(row, 7, f"=IF(B{row}=0,0,(C{row}+0.5*D{row})/B{row})")
        be.cell(row, 7).number_format = "0%"

        for col in range(1, 8):
            cell = be.cell(row, col)
            cell.fill = fill
            cell.border = THIN
            if col > 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")

    # Ensure column G width for Not Started / Progress %
    be.column_dimensions["G"].width = 14
    if not be.column_dimensions["A"].width:
        be.column_dimensions["A"].width = 52
    if not be.column_dimensions["E"].width or be.column_dimensions["E"].width < 40:
        be.column_dimensions["E"].width = 78

    # Subtitle note under % if A2 empty-ish — leave schedule notes alone
    if be["A2"].value is None:
        be["A2"] = (
            "Started first week of July. August commitment = DEVELOPMENT only (through 2026-08-31). "
            "Frontend integration TBD. Testing & documentation after development. "
            "Progress % below uses Backend Progress Tracker (same scoring as Frontend)."
        )
        be["A2"].font = Font(italic=True, color="FF666666", size=10)
        try:
            be.merge_cells("A2:F2")
        except Exception:
            pass

    wb.save(PATH)
    print(f"Updated: {PATH}")
    print(f"Tracker tasks: {len(TRACKER_ROWS)}")
    print("Sheets:", wb.sheetnames)

    # Quick sanity counts
    from collections import Counter

    c = Counter(s for _, _, s in TRACKER_ROWS)
    print("Status counts:", dict(c))
    earned = sum(1 if s == "Completed" else 0.5 if s == "In Progress" else 0 for _, _, s in TRACKER_ROWS)
    print(f"Expected overall ~ {round(100 * earned / len(TRACKER_ROWS))}%")


if __name__ == "__main__":
    main()
